"""Copia los valores calculados (por LibreOffice) dentro del .xlsx generado por openpyxl.

openpyxl escribe fórmulas sin valor en caché; Excel las recalcula al abrir, pero los visores
(previsualización web/móvil) mostrarían celdas vacías. Este paso mantiene intactas las fórmulas,
gráficos, tablas, validaciones y formato del archivo original y solo agrega el <v> de cada fórmula.

Uso: python inject_cached_values.py original.xlsx recalculado.xlsx salida.xlsx
"""
import re
import sys
import zipfile
from datetime import date, datetime, time
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

src, calc, out = sys.argv[1:4]
values = load_workbook(calc, data_only=True)
names = load_workbook(src, read_only=True).sheetnames

CELL = re.compile(r'<c r="([A-Z]+\d+)"((?: [a-zA-Z]+="[^"]*")*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', re.S)


def encode(v):
    if v is None:
        return ' t="str"', ""
    if isinstance(v, bool):
        return ' t="b"', "1" if v else "0"
    if isinstance(v, (datetime, date)):
        return "", repr(float(to_excel(v)))
    if isinstance(v, time):
        return "", repr((v.hour * 3600 + v.minute * 60 + v.second) / 86400)
    if isinstance(v, (int, float)):
        return "", repr(v)
    s = str(v)
    if s.startswith("#"):
        return ' t="e"', escape(s)
    return ' t="str"', escape(s)


zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
filled = 0
for item in zin.infolist():
    data = zin.read(item.filename)
    m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
    if m:
        ws = values[names[int(m.group(1)) - 1]]
        text = data.decode("utf-8")

        def sub(mt):
            global filled
            ref, attrs, f = mt.group(1), mt.group(2), mt.group(3)
            attrs = re.sub(r' t="[^"]*"', "", attrs)
            t, v = encode(ws[ref].value)
            filled += 1
            return f'<c r="{ref}"{attrs}{t}><f>{f}</f><v>{v}</v></c>'
        data = CELL.sub(sub, text).encode("utf-8")
    zout.writestr(item, data)
zout.close()
print("valores incrustados:", filled)
